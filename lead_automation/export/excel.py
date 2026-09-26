"""Excel export with openpyxl: a formatted *Leads* sheet (filters, frozen
header, hyperlinks, colour-coded email status) and a *Summary* sheet with
the cleaning report and per-source counts. A CSV twin is written too."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..models import OUTPUT_COLUMNS, OUTPUT_HEADERS

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(bold=True, color="FFFFFF")
STATUS_FILL = {
    "found": PatternFill("solid", fgColor="E2F0D9"),
    "generated (unverified)": PatternFill("solid", fgColor="FFF2CC"),
    "missing": PatternFill("solid", fgColor="F8CBAD"),
}
LINK_FONT = Font(color="0563C1", underline="single")
COLUMN_WIDTHS = {
    "name": 38, "category": 34, "email": 32, "email_status": 22, "website": 36,
    "linkedin": 40, "location": 30, "country": 14, "description": 70,
    "source": 12, "source_url": 60, "collected_at": 20,
}


def _write_leads_sheet(ws, df: pd.DataFrame) -> None:
    ws.title = "Leads"
    ws.append([OUTPUT_HEADERS[c] for c in OUTPUT_COLUMNS])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")

    link_cols = {OUTPUT_COLUMNS.index(c) + 1 for c in ("website", "linkedin", "source_url")}
    email_col = OUTPUT_COLUMNS.index("email") + 1
    status_col = OUTPUT_COLUMNS.index("email_status") + 1

    for row in df[OUTPUT_COLUMNS].itertuples(index=False):
        ws.append(list(row))
        r = ws.max_row
        for c in link_cols:
            cell = ws.cell(row=r, column=c)
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink, cell.font = cell.value, LINK_FONT
        email = ws.cell(row=r, column=email_col)
        if isinstance(email.value, str) and "@" in email.value:
            email.hyperlink, email.font = f"mailto:{email.value}", LINK_FONT
        status = ws.cell(row=r, column=status_col)
        if status.value in STATUS_FILL:
            status.fill = STATUS_FILL[status.value]

    for idx, col in enumerate(OUTPUT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = COLUMN_WIDTHS.get(col, 20)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


def _write_summary_sheet(ws, df: pd.DataFrame, report: dict) -> None:
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
    for key, value in report.items():
        ws.append([key.replace("_", " ").capitalize(), value])

    for title, series in (
        ("Leads per source", df["source"].value_counts()),
        ("Email status breakdown", df["email_status"].value_counts()),
    ):
        ws.append([])
        ws.append([title, ""])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        for label, count in series.items():
            ws.append([label, int(count)])

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 24


def write_excel(df: pd.DataFrame, report: dict, path: Path) -> Path:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _write_leads_sheet(wb.active, df)
    _write_summary_sheet(wb.create_sheet(), df, report)
    wb.save(path)

    csv_path = path.with_suffix(".csv")
    df[OUTPUT_COLUMNS].rename(columns=OUTPUT_HEADERS).to_csv(csv_path, index=False, encoding="utf-8-sig")
    log.info("Wrote %s and %s (%d rows)", path, csv_path.name, len(df))
    return path
